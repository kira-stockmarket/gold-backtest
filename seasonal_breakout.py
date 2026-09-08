import os
import yfinance as yf
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import LineChart, Reference

OUTPUT_DIR = "public"
os.makedirs(OUTPUT_DIR, exist_ok=True)

print("Fetching GC=F data from Yahoo Finance...")
gold = yf.download('GC=F', start='2000-01-01', progress=False)
if isinstance(gold.columns, pd.MultiIndex):
    gold.columns = gold.columns.droplevel(1)
gold = gold.dropna(subset=['Close', 'High', 'Low'])
gold.index = pd.to_datetime(gold.index).tz_localize(None)

print("Calculating Breakout Signals...")
gold['Month'] = gold.index.month
gold['Is_Aug_Dec'] = gold['Month'].isin([8, 12])
gold['5D_High'] = gold['Close'].shift(1).rolling(window=5).max()
gold['Signal'] = gold['Is_Aug_Dec'] & (gold['Close'] > gold['5D_High'])

trades = []
in_trade = False
exit_day = 0

prices = gold['Close'].values
dates = gold.index
signals = gold['Signal'].values
months = gold['Month'].values

print("Executing 3-Day Positional Trades...")
for i in range(5, len(gold)-3):
    if not in_trade:
        if signals[i]:
            entry_price = prices[i]
            exit_price = prices[i+3]
            ret = (exit_price / entry_price) - 1
            
            trades.append({
                'Entry_Date': dates[i].strftime('%Y-%m-%d'),
                'Month': 'August' if months[i] == 8 else 'December',
                'Entry_Price': entry_price,
                'Exit_Date': dates[i+3].strftime('%Y-%m-%d'),
                'Exit_Price': exit_price,
                'Return': ret,
                'Result': 'Win' if ret > 0 else 'Loss'
            })
            in_trade = True
            exit_day = i + 3
    else:
        if i >= exit_day:
            in_trade = False

df_trades = pd.DataFrame(trades)
df_trades['Cum_Return'] = df_trades['Return'].cumsum()

print("Generating Excel Dashboard...")
wb = Workbook()
ws_dash = wb.active
ws_dash.title = "Strategy Dashboard"
ws_data = wb.create_sheet(title="Trade Ledger")

headers = list(df_trades.columns)
ws_data.append(headers)
for r in dataframe_to_rows(df_trades, index=False, header=False):
    ws_data.append(r)

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
align_center = Alignment(horizontal="center")
thin_border = Border(bottom=Side(style='thin', color='DDDDDD'))

for col_num in range(1, len(headers)+1):
    cell = ws_data.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center
    
for row in ws_data.iter_rows(min_row=2, max_row=ws_data.max_row, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.border = thin_border
        cell.alignment = align_center
        if cell.column in [3, 5]: 
            cell.number_format = '#,##0.00'
        elif cell.column in [6, 8]: 
            cell.number_format = '0.00%'
        elif cell.column == 7:
            cell.font = Font(color="27AE60", bold=True) if cell.value == 'Win' else Font(color="C0392B", bold=True)

for col in ws_data.columns:
    ws_data.column_dimensions[col[0].column_letter].width = 16

ws_dash.append(["Gold Seasonal Breakout Strategy (Aug & Dec)"])
ws_dash.cell(row=1, column=1).font = Font(size=16, bold=True, color="1F4E78")
ws_dash.append([])

stats = [
    ["Total Trades Evaluated", len(df_trades)],
    ["Win Rate", f"{(df_trades['Return'] > 0).mean()*100:.1f}%"],
    ["Average Return per Trade", f"{df_trades['Return'].mean()*100:.2f}%"],
    ["Total Cumulative Return", f"{df_trades['Cum_Return'].iloc[-1]*100:.2f}%"]
]

for stat in stats:
    ws_dash.append(stat)

for r in range(3, 7):
    ws_dash.cell(row=r, column=1).font = Font(bold=True)
    ws_dash.cell(row=r, column=1).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws_dash.cell(row=r, column=2).alignment = Alignment(horizontal="right")
    ws_dash.column_dimensions['A'].width = 28
    ws_dash.column_dimensions['B'].width = 15

chart = LineChart()
chart.title = "Strategy Cumulative Return (3-Day Hold)"
chart.style = 2
chart.y_axis.title = 'Cumulative % Return'
chart.x_axis.title = 'Trade Number'
chart.width = 20
chart.height = 10

data = Reference(ws_data, min_col=8, min_row=1, max_row=ws_data.max_row)
chart.add_data(data, titles_from_data=True)
ws_dash.add_chart(chart, "D3")

excel_path = os.path.join(OUTPUT_DIR, "gold_seasonal_breakout.xlsx")
wb.save(excel_path)
print(f"✅ Excel file saved successfully to {excel_path}")
