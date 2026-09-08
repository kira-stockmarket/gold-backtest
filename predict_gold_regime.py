import os
import yfinance as yf
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')  # Headless mode for GitHub Actions
import matplotlib.pyplot as plt
from datetime import timedelta
from sklearn.cluster import KMeans

# Ensure openpyxl is installed for Excel generation
try:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("openpyxl is required. Please ensure it is in requirements.txt.")

# 1. SETUP & CONSTANTS
OUTPUT_DIR = "public"
os.makedirs(OUTPUT_DIR, exist_ok=True)
API_KEY = os.getenv("COMTRADE_API_KEY")

DIWALI_DATES = {
    2010: '2010-11-05', 2011: '2011-10-26', 2012: '2012-11-13',
    2013: '2013-11-03', 2014: '2014-10-23', 2015: '2015-11-11',
    2016: '2016-10-30', 2017: '2017-10-19', 2018: '2018-11-07',
    2019: '2019-10-27', 2020: '2020-11-14', 2021: '2021-11-04',
    2022: '2022-10-24', 2023: '2023-11-12', 2024: '2024-10-31',
    2025: '2025-10-21'
}

# 2. UN COMTRADE API FETCH (Fixed to include missing parameters)
def get_india_gold_imports():
    fallback_data = {
        2010: 958.0, 2011: 969.0, 2012: 860.0, 2013: 825.0, 2014: 890.0, 
        2015: 915.0, 2016: 557.0, 2017: 778.0, 2018: 759.0, 2019: 690.0, 
        2020: 430.0, 2021: 1050.0, 2022: 706.0, 2023: 781.0, 2024: 757.0, 
        2025: 721.0
    }
    
    if not API_KEY:
        print("⚠️ No API Key found in environment. Using baseline historical data.")
        return fallback_data

    try:
        import comtradeapicall
        print("🔄 Connecting to UN Comtrade API for live physical trade data...")
        years_str = "2015,2016,2017,2018,2019,2020,2021,2022,2023,2024"
        
        # Explicitly passing the previously missing parameters as None
        df = comtradeapicall.getFinalData(
            subscription_key=API_KEY,
            typeCode='C', freqCode='A', clCode='HS', period=years_str,
            reporterCode='356', cmdCode='7108', flowCode='M', partnerCode='0',
            partner2Code=None, customsCode=None, motCode=None
        )
        
        if df is None or df.empty:
            print("⚠️ Comtrade returned empty data. Using fallback.")
            return fallback_data
            
        api_data = {}
        for _, row in df.iterrows():
            year = int(row['period'])
            api_data[year] = float(row['netWgt']) / 1000.0 
            
        for y in fallback_data:
            if y not in api_data:
                api_data[y] = fallback_data[y]
                
        print("✅ Successfully pulled data from UN Comtrade API!")
        return api_data

    except Exception as e:
        print(f"❌ Comtrade API Error: {e}. Falling back to baseline data.")
        return fallback_data

def get_price_at(df, target_date):
    target = pd.to_datetime(target_date)
    sub = df.loc[df.index <= target, 'Close']
    if len(sub) == 0:
        return np.nan
    val = sub.iloc[-1]
    return float(val.iloc[0]) if isinstance(val, pd.Series) else float(val)

# 3. EXCEL EXPORT (Matches your requested format)
def generate_excel_report(df, cluster_summary):
    print("📊 Generating Excel ML Backtest Report...")
    excel_df = pd.DataFrame()
    excel_df['Event_ID'] = [f"EVT_{i+1:03d}" for i in range(len(df))]
    excel_df['Event_Name'] = df['Year'].apply(lambda y: f"Diwali_{y}")
    excel_df['T0_Date'] = pd.to_datetime(df['Diwali_Date']).dt.strftime('%d-%m-%Y')
    excel_df['Assigned_Cluster'] = df['Regime_Cluster']
    excel_df['Train_Clus'] = df['Regime_Cluster']  # Utilizing the K-Means cluster here
    
    excel_df['Feature_E'] = (df['Buildup_Return_%'] / 100).round(4)
    excel_df['Feature_P'] = (df['Pre_Diwali_15d_%'] / 100).round(4)
    
    def get_trend(val): return 'uptrend' if val > 0 else 'downtrend'

    excel_df['Pred_Pre_Trend'] = 'uptrend'
    excel_df['Actual_Pre_Ret'] = excel_df['Feature_P']
    excel_df['Actual_Pre_Trend'] = excel_df['Actual_Pre_Ret'].apply(get_trend)
    excel_df['Pre_Success'] = excel_df['Pred_Pre_Trend'] == excel_df['Actual_Pre_Trend']

    cluster_means = cluster_summary.set_index('Regime_Cluster')['Avg_Post15']
    excel_df['Pred_Post_Trend'] = df['Regime_Cluster'].map(lambda c: get_trend(cluster_means[c]))
    
    excel_df['Actual_Post_Ret'] = (df['Post_Diwali_15d_%'] / 100).round(4)
    excel_df['Actual_Post_Trend'] = excel_df['Actual_Post_Ret'].apply(get_trend)
    excel_df['Post_Success'] = excel_df['Pred_Post_Trend'] == excel_df['Actual_Post_Trend']

    wb = Workbook()
    ws = wb.active
    ws.title = "ML_Regime_Backtest"

    header_fill = PatternFill(start_color="34495E", end_color="34495E", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(left=Side(style='thin', color='BDC3C7'), right=Side(style='thin', color='BDC3C7'),
                         top=Side(style='thin', color='BDC3C7'), bottom=Side(style='thin', color='BDC3C7'))

    headers = ['Event_ID', 'Event_Name', 'T0_Date', 'Assigned_', 'Train_Clus', 'Feature_E', 'Feature_P', 
               'Pred_Pre_', 'Actual_Pre', 'Actual_Pre', 'Pre_Secto', 'Pred_Post', 'Actual_Po', 'Actual_Po', 'Post_Secto']
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill; cell.font = header_font; cell.alignment = align_center; cell.border = thin_border

    for r in dataframe_to_rows(excel_df, index=False, header=False):
        ws.append(r)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border; cell.alignment = align_center
            if cell.column in [6, 7, 9, 13]: cell.number_format = '0.0000'
            if cell.column in [11, 15]:
                cell.font = Font(color="27AE60", bold=True) if cell.value == True else Font(color="C0392B", bold=True)
            if cell.column in [8, 10, 12, 14]:
                if cell.value == 'uptrend': cell.font = Font(color="27AE60")
                elif cell.value == 'downtrend': cell.font = Font(color="C0392B")

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column].width = max_length + 3
    ws.freeze_panes = "A2"

    excel_path = os.path.join(OUTPUT_DIR, "gold_ml_backtest.xlsx")
    wb.save(excel_path)
    print(f"✅ Excel saved to {excel_path}")

# 4. CORE ENGINE & HTML GENERATOR
def run_pipeline():
    print("📈 Fetching historical Gold Futures (GC=F)...")
    gold = yf.download('GC=F', start='2009-01-01', progress=False)
    if isinstance(gold.columns, pd.MultiIndex):
        gold.columns = gold.columns.droplevel(1)
    gold.index = pd.to_datetime(gold.index).tz_localize(None)

    imports_data = get_india_gold_imports()
    records = []
    
    print("🧠 Extracting patterns & running clustering algorithm...")
    for year, diwali_str in DIWALI_DATES.items():
        diwali_date = pd.to_datetime(diwali_str)
        t_minus_45 = diwali_date - timedelta(days=45)
        t_minus_15 = diwali_date - timedelta(days=15)
        t_post_15 = diwali_date + timedelta(days=15)

        p_45 = get_price_at(gold, t_minus_45)
        p_15 = get_price_at(gold, t_minus_15)
        p_diwali = get_price_at(gold, diwali_date)
        p_post = get_price_at(gold, t_post_15)

        if any(pd.isna(x) for x in [p_45, p_15, p_diwali, p_post]):
            continue

        annual_imports = imports_data.get(year, 700.0) 
        build_up_return = ((p_15 - p_45) / p_45) * 100
        pre_event_return = ((p_diwali - p_15) / p_15) * 100
        post_event_return = ((p_post - p_diwali) / p_diwali) * 100

        records.append({
            'Year': year,
            'Diwali_Date': diwali_str,
            'Imports_Tons': round(annual_imports, 1),
            'Buildup_Return_%': round(build_up_return, 2),
            'Pre_Diwali_15d_%': round(pre_event_return, 2),
            'Post_Diwali_15d_%': round(post_event_return, 2)
        })

    df = pd.DataFrame(records)
    X = df[['Buildup_Return_%', 'Pre_Diwali_15d_%', 'Imports_Tons']].values
    kmeans = KMeans(n_clusters=3, random_state=42, n_init=10)
    df['Regime_Cluster'] = kmeans.fit_predict(X)

    cluster_summary = df.groupby('Regime_Cluster').agg(
        Count=('Year', 'count'),
        Avg_Pre15=('Pre_Diwali_15d_%', 'mean'),
        Avg_Post15=('Post_Diwali_15d_%', 'mean'),
        Post_WinRate=('Post_Diwali_15d_%', lambda x: (x > 0).mean() * 100)
    ).reset_index()

    def name_regime(row):
        if row['Avg_Pre15'] > 1.0 and row['Avg_Post15'] < 0: return "Exhaustion Spike (Sell-the-Fact)"
        elif row['Avg_Post15'] > 0.5: return "Sustained Rally (Follow-through)"
        else: return "Macro Chop (Consolidation)"

    cluster_summary['Regime_Name'] = cluster_summary.apply(name_regime, axis=1)
    df = df.merge(cluster_summary[['Regime_Cluster', 'Regime_Name']], on='Regime_Cluster', how='left')

    latest = df.iloc[-1]
    latest_regime = cluster_summary[cluster_summary['Regime_Cluster'] == latest['Regime_Cluster']].iloc[0]

    # Explicitly define plot path and save figure
    plot_path = os.path.join(OUTPUT_DIR, "predictive_clusters.png")
    fig, ax = plt.subplots(figsize=(10, 5))
    colors = ['#3498db', '#e74c3c', '#2ecc71']
    for c in range(3):
        sub = df[df['Regime_Cluster'] == c]
        ax.scatter(sub['Pre_Diwali_15d_%'], sub['Post_Diwali_15d_%'], 
                   s=sub['Imports_Tons'], c=colors[c], alpha=0.7, 
                   label=f"{cluster_summary.loc[c, 'Regime_Name']} (Win: {cluster_summary.loc[c, 'Post_WinRate']:.0f}%)")

    ax.axhline(0, color='black', linestyle='--', alpha=0.5)
    ax.axvline(0, color='black', linestyle='--', alpha=0.5)
    ax.set_xlabel("Pre-Diwali 15-Day Return (%)", fontweight='bold')
    ax.set_ylabel("Post-Diwali 15-Day Return (%) [TARGET]", fontweight='bold')
    ax.set_title("Gold Regimes: Physical Volume (Bubble Size) vs Momentum")
    ax.legend(loc='best')
    plt.tight_layout()
    plt.savefig(plot_path)
    plt.close()

    # Generate outputs
    generate_excel_report(df, cluster_summary)
    generate_html(df, cluster_summary, latest, latest_regime)

def generate_html(df, summary, latest, latest_regime):
    html_path = os.path.join(OUTPUT_DIR, "index.html")
    
    summary_table = summary[['Regime_Name', 'Count', 'Avg_Pre15', 'Avg_Post15', 'Post_WinRate']].to_html(
        index=False, classes="table", float_format=lambda x: f"{x:.2f}"
    )
    history_table = df[['Year', 'Diwali_Date', 'Imports_Tons', 'Pre_Diwali_15d_%', 'Post_Diwali_15d_%', 'Regime_Name']].to_html(
        index=False, classes="table"
    )

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Gold Seasonal ML Prediction Engine</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, sans-serif; margin: 30px auto; max-width: 900px; padding: 0 20px; color: #1e293b; background: #f8fafc; }}
        .prediction-card {{ background: #fff; border: 2px solid #3b82f6; border-radius: 12px; padding: 20px; box-shadow: 0 4px 6px -1px rgba(0,0,0,0.05); }}
        .badge {{ background: #dbeafe; color: #1d4ed8; font-weight: bold; padding: 4px 10px; border-radius: 9999px; font-size: 0.85rem; }}
        .stats-grid {{ display: grid; grid-template-columns: repeat(3, 1fr); gap: 16px; margin: 16px 0; }}
        .stat-box {{ background: #f1f5f9; padding: 14px; border-radius: 8px; text-align: center; }}
        .stat-val {{ font-size: 1.5rem; font-weight: bold; margin-top: 6px; }}
        .green {{ color: #16a34a; }} .red {{ color: #dc2626; }}
        img {{ width: 100%; border-radius: 10px; border: 1px solid #cbd5e1; margin: 20px 0; }}
        table {{ width: 100%; border-collapse: collapse; background: #fff; border-radius: 8px; overflow: hidden; margin-top: 15px; border: 1px solid #e2e8f0; }}
        th, td {{ padding: 10px; text-align: left; border-bottom: 1px solid #e2e8f0; }}
        th {{ background: #f8fafc; color: #475569; }}
        .dl-btn {{ display: inline-block; background: #27ae60; color: #fff; padding: 10px 15px; text-decoration: none; border-radius: 5px; font-weight: bold; margin-bottom: 15px; }}
    </style>
</head>
<body>
    <h1>Gold Backtester (Comtrade + Yahoo Finance)</h1>
    <a href="gold_ml_backtest.xlsx" download class="dl-btn">📥 Download ML Excel Data</a>
    
    <div class="prediction-card">
        <span class="badge">Live ML Prediction Engine</span>
        <h2>Latest Setup ({latest['Year']}): {latest_regime['Regime_Name']}</h2>
        <p>Based on India importing {latest['Imports_Tons']} Tons of physical gold and pre-festival momentum of {latest['Pre_Diwali_15d_%']:+.2f}%, the KMeans algorithm has classified the current market into the regime above.</p>
        <div class="stats-grid">
            <div class="stat-box"><div>Historical Win Rate</div><div class="stat-val green">{latest_regime['Post_WinRate']:.1f}%</div></div>
            <div class="stat-box"><div>Expected Forward Return</div><div class="stat-val {'green' if latest_regime['Avg_Post15'] > 0 else 'red'}">{latest_regime['Avg_Post15']:+.2f}%</div></div>
            <div class="stat-box"><div>Historical Occurrences</div><div class="stat-val">{int(latest_regime['Count'])}</div></div>
        </div>
    </div>

    <h2>Regime Clustering Analysis</h2>
    <img src="predictive_clusters.png" alt="Cluster Plot">
    <h2>Regime Profiles</h2>
    {summary_table}
    <h2>Full Ledger</h2>
    {history_table}
</body>
</html>"""

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"✅ Generated final website at {html_path}")

if __name__ == "__main__":
    run_pipeline()
